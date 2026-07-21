import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from src.config import company_config as cfg


def _argb(hex_color: str) -> str:
    """openpyxl wants an 8-digit ARGB hex string; our config stores plain '#rrggbb'."""
    return "FF" + hex_color.lstrip("#").upper()


class ExcelExporter:
    @staticmethod
    def export_sales_ledger(records: list, output_path: str) -> bool:
        """
        Saves a structured, beautifully stylized Excel sheet of transactions,
        with numeric value masks and dynamic bottom SUM formula columns.

        NOTE: `records` is expected to carry the EFFECTIVE status/due (i.e. already
        accounting for any linked due-payments) — history_controller._get_date_range_records
        supplies invoice.effective_status / invoice.effective_due_amount, so this exporter
        doesn't need to know anything about the due-payment chain itself.
        """
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sales History Logs"

            ws.views.sheetView[0].showGridLines = True

            font_title = Font(name='Segoe UI', size=16, bold=True, color='2C3E50')
            font_header = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
            font_body = Font(name='Segoe UI', size=11, bold=False)
            font_total = Font(name='Segoe UI', size=11, bold=True)

            fill_header = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
            fill_total = PatternFill(start_color='ECF0F1', end_color='ECF0F1', fill_type='solid')

            align_left = Alignment(horizontal='left', vertical='center')
            align_right = Alignment(horizontal='right', vertical='center')
            align_center = Alignment(horizontal='center', vertical='center')

            border_thin = Side(border_style="thin", color="CBD5E1")
            border_double = Side(border_style="double", color="2C3E50")

            cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
            total_border = Border(top=border_thin, bottom=border_double)

            # 1. Sheet Title
            ws.merge_cells('A1:G1')
            ws['A1'] = f"{cfg.COMPANY_NAME.upper()} - SALES HISTORY LEDGER"
            ws['A1'].font = font_title
            ws['A1'].alignment = align_left
            ws.row_dimensions[1].height = 35

            ws.append([])

            # 2. Table Headers  (mirrors the on-screen History Ledger column order)
            headers = ["Invoice ID", "Customer Name", "Transaction Date", "Discount", "Grand Total", "Status", "Due"]
            ws.append(headers)
            header_row_idx = 3
            ws.row_dimensions[header_row_idx].height = 25

            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=header_row_idx, column=col_idx)
                cell.font = font_header
                cell.fill = fill_header
                cell.alignment = align_center if col_idx != 2 else align_left
                cell.border = cell_border

            # 3. Append Ledger Records
            start_row_data = 4
            for index, row in enumerate(records, start=start_row_data):
                status_text = row.get('status', '')
                ws.append([
                    f"#{row['id']}",
                    row['customer_name'],
                    row['date'],
                    row['discount'],
                    row['grand_total'],
                    status_text,
                    row['due']
                ])

                ws.row_dimensions[index].height = 20
                for col_idx in range(1, 8):
                    cell = ws.cell(row=index, column=col_idx)
                    cell.font = font_body
                    cell.border = cell_border

                    if col_idx == 1 or col_idx == 3:
                        cell.alignment = align_center
                    elif col_idx in (4, 5, 7):
                        cell.alignment = align_right
                        cell.number_format = f'"{cfg.CURRENCY_SYMBOL}"#,##0.00'
                    elif col_idx == 6:
                        # Status column: color-coded + bold, same resolver the UI/PDF use
                        cell.alignment = align_center
                        cell.font = Font(name='Segoe UI', size=11, bold=True,
                                          color=_argb(cfg.resolve_status_color(status_text)))

            # 4. Total Calculation Row
            end_row_data = ws.max_row
            total_row_idx = end_row_data + 1
            ws.row_dimensions[total_row_idx].height = 22

            ws.cell(row=total_row_idx, column=1, value="").border = total_border
            ws.cell(row=total_row_idx, column=2, value="").border = total_border
            ws.cell(row=total_row_idx, column=3, value="Totals:").font = font_total
            ws.cell(row=total_row_idx, column=3).alignment = align_right
            ws.cell(row=total_row_idx, column=3).border = total_border

            discount_formula_cell = ws.cell(row=total_row_idx, column=4, value=f"=SUM(D{start_row_data}:D{end_row_data})")
            discount_formula_cell.font = font_total
            discount_formula_cell.alignment = align_right
            discount_formula_cell.number_format = f'"{cfg.CURRENCY_SYMBOL}"#,##0.00'
            discount_formula_cell.border = total_border

            revenue_formula_cell = ws.cell(row=total_row_idx, column=5, value=f"=SUM(E{start_row_data}:E{end_row_data})")
            revenue_formula_cell.font = font_total
            revenue_formula_cell.alignment = align_right
            revenue_formula_cell.number_format = f'"{cfg.CURRENCY_SYMBOL}"#,##0.00'
            revenue_formula_cell.border = total_border

            ws.cell(row=total_row_idx, column=6, value="").border = total_border

            due_formula_cell = ws.cell(row=total_row_idx, column=7, value=f"=SUM(G{start_row_data}:G{end_row_data})")
            due_formula_cell.font = font_total
            due_formula_cell.alignment = align_right
            due_formula_cell.number_format = f'"{cfg.CURRENCY_SYMBOL}"#,##0.00'
            due_formula_cell.border = total_border

            for col_idx in range(1, 8):
                ws.cell(row=total_row_idx, column=col_idx).fill = fill_total

            # 5. Auto-adjust columns
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row == total_row_idx and col_letter in ('D', 'E', 'G'):
                        continue
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

            wb.save(output_path)
            return True

        except Exception as e:
            print(f"[Critical Excel Exporter Error]: {e}")
            return False